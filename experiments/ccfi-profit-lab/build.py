"""Build an offline, read-only COSCO earnings scenario page."""
import argparse
import hashlib
import json
import sqlite3
import sys
from pathlib import Path

import numpy as np
import openpyxl

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
from stock_analysis.domains.financials.financial_storage import decompress_json


def shifted_months(year, quarter, lag):
    start = year * 12 + (quarter - 1) * 3 - lag
    return [f'{i // 12:04d}-{i % 12 + 1:02d}' for i in range(start, start + 3)]


def rate(months, year, quarter, lag):
    keys = shifted_months(year, quarter, lag)
    values = [months.get(k) for k in keys]
    return None if any(v is None for v in values) else sum(values) / 3


def fit(rows):
    x = np.array([r['ccfi'] for r in rows], dtype=float)
    y = np.array([r['profit'] for r in rows], dtype=float)
    slope, intercept = np.linalg.lstsq(np.column_stack([x, np.ones(len(x))]), y, rcond=None)[0]
    residuals = y - (slope*x+intercept)
    return dict(slope=float(slope), intercept=float(intercept), n=len(x),
                r2=float(1 - sum(residuals**2)/sum((y-y.mean())**2)),
                rmse=float(np.sqrt(np.mean(residuals**2))))


def load_financials(db_path):
    conn = sqlite3.connect(f'file:{db_path.as_posix()}?mode=ro', uri=True)
    conn.row_factory = sqlite3.Row
    identity = conn.execute("SELECT id FROM instruments WHERE exchange='sh' AND code='601919'").fetchone()[0]
    raw = {}
    for row in conn.execute("SELECT * FROM financial_reports WHERE instrument_id=? AND period BETWEEN '2022-12-31' AND '2026-06-30' AND report_type IN ('gjzb','lrb') ORDER BY period,source", (identity,)):
        obj = json.loads(decompress_json(row['raw_json']))
        raw[(row['period'],row['report_type'])] = (obj['data'], dict(source=row['source'], hash=row['content_hash'], obtained=row['obtained_at'], published=row['publish_date']))
    def metric(period, field, gj=True):
        for kind in (['gjzb','lrb'] if gj else ['lrb']):
            items, provenance = raw.get((period,kind), ([],None))
            for i in items:
                if i.get('item_field') == field and i.get('item_value') not in [None,'','--']:
                    return float(i['item_value'])/1e8, dict(**provenance, table='financial_reports', period=period, report_type=kind, field=field, unit='元')
        return None, None
    records = []
    ends = ['03-31','06-30','09-30','12-31']
    for year in range(2023,2027):
        for q in range(1,5):
            period = f'{year}-{ends[q-1]}'
            if period > '2026-06-30': continue
            values, sources = {}, {}
            for key,field,gj in [('profit','PARENETP',True),('revenue','BIZTOTINCO',True),('cost','BIZCOST',False),('deducted','NPCUT',True),('pretax','TOTPROFIT',False),('net','NETPROFIT',True)]:
                cumulative, source = metric(period,field,gj)
                prev, ps = metric(f'{year}-{ends[q-2]}',field,gj) if q>1 else (0,None)
                values[key] = cumulative-prev if cumulative is not None and prev is not None else None
                sources[key] = dict(current=source, previous=ps)
            if values['profit'] is not None:
                records.append(dict(period=period,label=f'{year}Q{q}',year=year,q=q,**values,sources=sources))
    conn.close()
    return records


def build(workbook, database, reference):
    wb = openpyxl.load_workbook(workbook,read_only=True,data_only=True)
    months, monthly, weekly = {}, [], []
    for idx,row in enumerate(wb['月度CCFI'].iter_rows(min_row=6,max_row=17,values_only=True),6):
        month = int(str(row[0]).replace('月',''))
        for col,year in enumerate(range(2023,2027),1):
            value = row[col]
            key=f'{year}-{month:02d}'
            months[key]=value
            monthly.append(dict(month=key,value=value,note=row[5],cell=f'{chr(65+col)}{idx}'))
    for row in wb['每周CCFI'].iter_rows(min_row=6,max_row=57,values_only=True):
        weekly.append(dict(week=row[0],values=list(row[1:5]),note=row[5]))
    history=load_financials(database)
    common=[r for r in history if all(rate(months,r['year'],r['q'],lag) is not None for lag in [0,1,2])]
    models={}
    for lag in [0,1,2]:
        rows=[dict(**r,ccfi=rate(months,r['year'],r['q'],lag)) for r in common]
        model=fit(rows)
        backtests=[]
        for i in range(6,len(rows)):
            training=rows[:i]
            f=fit(training)
            previous=rows[i-1]
            pred=previous['profit']+f['slope']*(rows[i]['ccfi']-previous['ccfi'])
            backtests.append(dict(label=rows[i]['label'],actual=rows[i]['profit'],prediction=pred,error=pred-rows[i]['profit'],baseline=previous['profit'],training_end=previous['label']))
        model.update(rows=rows,backtests=backtests,mae=float(np.mean([abs(r['error']) for r in backtests])),baseline_mae=float(np.mean([abs(r['baseline']-r['actual']) for r in backtests])),target=rate(months,2026,3,lag),anchor=rate(months,2026,2,lag))
        models[str(lag)]=model
    reports=[]
    for p in sorted((ROOT/'data/company_reports/sh601919').glob('*/v1/manifest.json')):
        m=json.loads(p.read_text(encoding='utf-8'))
        reports.append(dict(period=m['report_period'],url=m['source_url'],hash=m['content_hash'],path=str(p.with_name('report.pdf').relative_to(ROOT)).replace('\\','/')))
    q2=history[-1]
    # Business amounts are in RMB thousand in the reports; convert to RMB 100m.
    international=((97479277-6775664)-(45092960-3388811))/1e5
    domestic=(6775664-3388811)/1e5
    for label in ['2026Q1','2024FY']:
        sample=ROOT/f'data/verification/samples/ccfi-profit-lab/{label}.source.json'
        if sample.exists():
            m=json.loads(sample.read_text(encoding='utf-8'))
            reports.append(dict(period=label,url=m['url'],hash=m['sha256'],path=f'data/verification/samples/ccfi-profit-lab/{label}.pdf'))
    data=dict(asof='2026-10-24', cutoff='2026-06-30', monthly=monthly,weekly=weekly,months=months,history=history,models=models,reports=reports,
              workbook=dict(name=workbook.name,hash=hashlib.sha256(workbook.read_bytes()).hexdigest()),
              reference=dict(url='https://xueqiu.com/5074021865/409286919',author='ljm94',published='2026-09-15',hash=hashlib.sha256(reference.read_bytes()).hexdigest() if reference.exists() else None),
              business=dict(international=international,domestic=domestic,other=q2['revenue']-international-domestic,cost=q2['cost'],teu=(14279451-3152616)-(6915550-1518912),below=q2['pretax']-(q2['revenue']-q2['cost']),conversion=q2['profit']/q2['pretax']))
    calibration_path=ROOT/'data/verification/reports/ccfi-profit-lab/calibration.json'
    if calibration_path.exists():
        calibration=json.loads(calibration_path.read_text(encoding='utf-8'))
        data['calibration']={k:v for k,v in calibration.items() if k!='extraction'}
        data['calibration']['reports']=json.loads((ROOT/'data/verification/samples/ccfi-profit-lab/calibration_manifest.json').read_text(encoding='utf-8'))
    weekly_path=ROOT/'data/verification/reports/ccfi-profit-lab/weekly.json'
    if weekly_path.exists():
        weekly_study=json.loads(weekly_path.read_text(encoding='utf-8'))
        if weekly_study['workbook']['hash']!=data['workbook']['hash']:
            raise ValueError('周度研究的Excel版本已过期，请重新计算weekly.py')
        data['weeklyStudy']=weekly_study
    audit=ROOT/'data/verification/reports/ccfi-profit-lab'
    audit.mkdir(parents=True,exist_ok=True)
    (audit/'snapshot.json').write_text(json.dumps(data,ensure_ascii=False,indent=2,allow_nan=False),encoding='utf-8')
    html=(Path(__file__).with_name('template.html')).read_text(encoding='utf-8').replace('__DATA__',json.dumps(data,ensure_ascii=False,allow_nan=False).replace('</','<\\/')).replace('__CALIBRATION_SCRIPT__',Path(__file__).with_name('calibration.js').read_text(encoding='utf-8')).replace('__WEEKLY_SECTION__',Path(__file__).with_name('weekly_section.html').read_text(encoding='utf-8')).replace('__WEEKLY_SCRIPT__',Path(__file__).with_name('weekly_view.js').read_text(encoding='utf-8'))
    Path(__file__).with_name('index.html').write_text(html,encoding='utf-8')
    print(json.dumps({k:{f:v for f,v in m.items() if f not in ['rows','backtests']} for k,m in models.items()},ensure_ascii=False,indent=2))


if __name__=='__main__':
    parser=argparse.ArgumentParser()
    parser.add_argument('--workbook',type=Path,required=True)
    parser.add_argument('--database',type=Path,default=ROOT/'data/stock_analysis.sqlite3')
    parser.add_argument('--reference-pdf',type=Path,required=True)
    args=parser.parse_args()
    build(args.workbook,args.database,args.reference_pdf)
